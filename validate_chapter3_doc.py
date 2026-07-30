from pathlib import Path
import sys

from docx import Document
from docx.oxml.ns import qn


def paragraph_num(paragraph):
    num_pr = paragraph._p.find("./" + qn("w:pPr") + "/" + qn("w:numPr"))
    if num_pr is None:
        return None, None
    num_id = num_pr.find(qn("w:numId"))
    ilvl = num_pr.find(qn("w:ilvl"))
    return (
        num_id.get(qn("w:val")) if num_id is not None else None,
        ilvl.get(qn("w:val")) if ilvl is not None else None,
    )


def first_line_twips(paragraph):
    indent = paragraph._p.find("./" + qn("w:pPr") + "/" + qn("w:ind"))
    return indent.get(qn("w:firstLine")) if indent is not None else None


def section_text(doc, start_text):
    start = next(
        index
        for index, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip() == start_text
    )
    return [paragraph.text for paragraph in doc.paragraphs[start:]]


def border_for(borders, *alternatives):
    return next(
        (
            borders.find(qn(f"w:{edge}"))
            for edge in alternatives
            if borders.find(qn(f"w:{edge}")) is not None
        ),
        None,
    )


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: validate_chapter3_doc.py SOURCE.docx OUTPUT.docx")
    source_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    source = Document(source_path)
    doc = Document(output_path)

    start = next(
        index
        for index, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip() == "功能设计"
    )
    end = next(
        index
        for index, paragraph in enumerate(doc.paragraphs)
        if index > start and paragraph.text.strip() == "系统实现"
    )
    chapter = doc.paragraphs[start:end]

    expected_h2 = [
        "网页端双层创作与参数调节功能",
        "照片纸面化与实体留存功能",
        "AI辅助创作与再创作功能",
        "小智语音语义控制功能",
    ]
    actual_h2 = [p.text for p in chapter if p.style.style_id == "2"]
    assert actual_h2 == expected_h2, actual_h2
    h3 = [p for p in chapter if p.style.style_id == "3"]
    assert len(h3) == 11, len(h3)

    for paragraph in chapter:
        if paragraph.style.style_id == "1":
            assert paragraph_num(paragraph) == ("44", "0")
        elif paragraph.style.style_id == "2":
            assert paragraph_num(paragraph) == ("44", "1")
        elif paragraph.style.style_id == "3":
            assert paragraph_num(paragraph) == ("44", "2")

    body_paragraphs = [
        p
        for p in chapter
        if p.style.style_id == "a"
        and p.text.strip()
        and not p.text.startswith("图3-")
        and not p.text.startswith("表3-")
    ]
    bad_indents = [
        (p.text[:24], first_line_twips(p))
        for p in body_paragraphs
        if first_line_twips(p) != "480"
    ]
    assert not bad_indents, bad_indents

    figure_captions = [p for p in chapter if p.text.startswith("图3-")]
    assert [p.text[:4] for p in figure_captions] == ["图3-1", "图3-2", "图3-3"]
    for paragraph in figure_captions:
        assert paragraph.alignment is not None
        for run in paragraph.runs:
            if run.text:
                r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
                assert r_fonts.get(qn("w:eastAsia")) == "楷体"
                assert run.font.size is not None and run.font.size.pt == 11.0

    table_captions = [p.text for p in chapter if p.text.startswith("表3-")]
    assert table_captions == [
        "表3-1 网页端主要功能参数及其作用",
        "表3-2 小智语音意图与功能联动示例",
    ]
    assert any("模糊或平滑强度" in p.text and "最短轨迹" in p.text for p in chapter)
    assert any("展示、装裱、赠送与长期留存" == p.text for p in chapter)
    assert any("普通照片增加新的创作可能" in p.text for p in chapter)
    assert any("同一操作无需依赖固定口令" in p.text for p in chapter)
    assert any("结构化命令层" in p.text and "设备状态" in p.text for p in chapter)
    assert sum("图片待替换" in p.text for p in figure_captions) == 3

    assert len(doc.sections) == len(source.sections) == 9
    assert len(doc.tables) == len(source.tables) + 4
    assert len(doc.inline_shapes) == len(source.inline_shapes) + 6
    assert section_text(doc, "系统实现") == section_text(source, "系统实现")

    for table in doc.tables[-4:]:
        borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")
        assert borders is not None
        for alternatives in (
            ("top",),
            ("start", "left"),
            ("bottom",),
            ("end", "right"),
            ("insideH",),
            ("insideV",),
        ):
            border = border_for(borders, *alternatives)
            assert border is not None and border.get(qn("w:val")) == "nil"

    print(
        f"OK paragraphs={len(doc.paragraphs)} chapter3_paragraphs={len(chapter)} "
        f"h2={len(actual_h2)} h3={len(h3)} body={len(body_paragraphs)} "
        f"captions={len(figure_captions)} tables={len(doc.tables)} "
        f"pictures={len(doc.inline_shapes)} sections={len(doc.sections)}"
    )


if __name__ == "__main__":
    main()
