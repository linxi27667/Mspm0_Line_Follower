from copy import deepcopy
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


CHAPTER_START = "设计需求分析"
CHAPTER_END = "特色与创新"


def find_paragraph(doc, text, start=0):
    for paragraph in doc.paragraphs[start:]:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def first_run_properties(paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return run._r.rPr
    return None


def paragraph_element_from_template(template_p, text, run_rpr=None):
    paragraph = deepcopy(template_p)
    for child in list(paragraph):
        if child.tag != qn("w:pPr"):
            paragraph.remove(child)

    run = OxmlElement("w:r")
    if run_rpr is not None:
        run.append(deepcopy(run_rpr))
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    paragraph.append(run)
    return paragraph


def page_break_element_from_template(template_p, run_rpr=None):
    paragraph = deepcopy(template_p)
    for child in list(paragraph):
        if child.tag != qn("w:pPr"):
            paragraph.remove(child)
    run = OxmlElement("w:r")
    if run_rpr is not None:
        run.append(deepcopy(run_rpr))
    page_break = OxmlElement("w:br")
    page_break.set(qn("w:type"), "page")
    run.append(page_break)
    paragraph.append(run)
    return paragraph


def set_run_font(run, east_asia, size, bold=None, color=None, italic=None):
    run.font.name = east_asia
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_margins(cell, top=65, start=100, bottom=65, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = None
    total_width = sum(widths)

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total_width))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "nil")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.0, keep_next=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if keep_next is not None:
        paragraph.paragraph_format.keep_with_next = keep_next


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def add_centered_image(cell, image_path, width_inches, alt_text, placeholder_note):
    paragraph = cell.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, after=1, keep_next=True)
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Inches(width_inches))
    inline._inline.docPr.set("descr", alt_text)

    note = cell.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(note, after=0)
    note_run = note.add_run(placeholder_note)
    set_run_font(note_run, "楷体", 9.5)


def add_labeled_paragraph(cell, label, text, first=False):
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    if first:
        clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_paragraph_spacing(paragraph, after=2, line=1.05)
    label_run = paragraph.add_run(label)
    set_run_font(label_run, "黑体", 10, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, "宋体", 10)


def build_comparison_table(doc, image_path):
    table = doc.add_table(rows=3, cols=2)
    set_table_geometry(table, [4415, 4415])

    titles = ("典型笛卡尔式XY写字机", "本作品：SCARA智能写画机")
    for index, title in enumerate(titles):
        cell = table.cell(0, index)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(paragraph, after=2, keep_next=True)
        run = paragraph.add_run(title)
        set_run_font(run, "黑体", 11, bold=True)

    add_centered_image(
        table.cell(1, 0),
        image_path,
        1.65,
        "占位图：典型笛卡尔式XY写字机",
        "（a）典型笛卡尔式XY写字机（待替换）",
    )
    add_centered_image(
        table.cell(1, 1),
        image_path,
        1.65,
        "占位图：本作品SCARA智能写画机",
        "（b）本作品SCARA智能写画机（待替换）",
    )

    left = table.cell(2, 0)
    right = table.cell(2, 1)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    add_labeled_paragraph(left, "优势：", "直角坐标关系直观，控制方法成熟，矩形工作区易规划。", first=True)
    add_labeled_paragraph(left, "局限：", "通常需要长行程导轨、横梁和外围框架，幅面增大时用料、体积与装配量同步增加。")
    add_labeled_paragraph(left, "交互：", "部分产品依赖专用上位机或离线文件，照片处理、排版与设备控制相互分离。")

    add_labeled_paragraph(right, "结构创新：", "双旋转臂覆盖书写区域，减少长直线模组，为小型化、便携化和低成本提供基础。", first=True)
    add_labeled_paragraph(right, "算法取舍：", "以逆运动学、轨迹规划、标定和误差补偿换取机械结构简化。")
    add_labeled_paragraph(right, "交互创新：", "网页端连接手机照片、AI处理、轨迹预览与写画输出。")
    return table


def build_workflow_table(doc, image_path):
    widths = [2600, 515, 2600, 515, 2600]
    table = doc.add_table(rows=2, cols=5)
    set_table_geometry(table, widths)

    image_specs = (
        (0, "占位图：手机拍摄或相册原图", "（a）手机拍摄或相册原图（占位）", "拍照或选择照片"),
        (2, "占位图：网页端AI处理与轨迹预览", "（b）网页端AI处理与轨迹预览（占位）", "AI优化、排版与轨迹预览"),
        (4, "占位图：写画机纸面输出成品", "（c）写画机纸面输出成品（占位）", "任务下发与自动书写"),
    )
    for column, alt_text, subcaption, label in image_specs:
        add_centered_image(
            table.cell(0, column),
            image_path,
            1.0,
            alt_text,
            subcaption,
        )
        label_cell = table.cell(1, column)
        label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = label_cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(paragraph, after=0, line=1.0)
        run = paragraph.add_run(label)
        set_run_font(run, "宋体", 9.5, bold=True)

    for column in (1, 3):
        for row in (0, 1):
            cell = table.cell(row, column)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(paragraph, after=0)
            if row == 0:
                run = paragraph.add_run("→")
                set_run_font(run, "宋体", 14, bold=True)
    return table


def main():
    if len(sys.argv) != 5:
        raise SystemExit(
            "Usage: revise_chapter1_innovation.py SOURCE.docx TEMPLATE.docx IMAGE.jpg OUTPUT.docx"
        )

    source_path = Path(sys.argv[1])
    template_path = Path(sys.argv[2])
    image_path = Path(sys.argv[3])
    output_path = Path(sys.argv[4])
    if not image_path.exists():
        raise FileNotFoundError(image_path)

    doc = Document(source_path)
    original_template = Document(template_path)
    chapter_start = find_paragraph(doc, CHAPTER_START)
    chapter_end = find_paragraph(doc, CHAPTER_END)

    start_index = next(
        index
        for index, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip() == CHAPTER_START
    )
    h2_template = next(
        p for p in doc.paragraphs[start_index + 1 :] if p.style.style_id == "2"
    )
    h3_template = next(
        p for p in doc.paragraphs[start_index + 1 :] if p.style.style_id == "3"
    )
    body_template = doc.paragraphs[26]
    figure_caption_template = find_paragraph(
        original_template, "图1-1 云计算三种服务模式的服务对比"
    )
    table_caption_template = find_paragraph(
        original_template, "表1-1云计算同传统IT服务模式的区别"
    )

    body = doc.element.body
    children = list(body)
    body_start = children.index(chapter_start._p)
    body_end = children.index(chapter_end._p)
    between = children[body_start + 1 : body_end]
    section_break = next(
        (
            element
            for element in reversed(between)
            if element.tag == qn("w:p")
            and element.find("./" + qn("w:pPr") + "/" + qn("w:sectPr"))
            is not None
        ),
        None,
    )
    if section_break is None:
        raise ValueError("Expected a section-break paragraph before the next chapter")
    for element in between:
        if element is not section_break:
            body.remove(element)

    templates = {
        "h2": (h2_template._p, first_run_properties(h2_template)),
        "h3": (h3_template._p, first_run_properties(h3_template)),
        "body": (body_template._p, first_run_properties(body_template)),
        "figure_caption": (
            figure_caption_template._p,
            first_run_properties(figure_caption_template),
        ),
        "table_caption": (
            table_caption_template._p,
            first_run_properties(table_caption_template),
        ),
    }

    def insert_paragraph(kind, text):
        template_p, run_rpr = templates[kind]
        element = paragraph_element_from_template(template_p, text, run_rpr)
        section_break.addprevious(element)

    insert_paragraph(
        "body",
        "本章围绕作品的设计需求、结构创新和使用特色进行概述。核心思路是以更复杂的算法和更严格的结构标定换取机械硬件简化，并通过网页端打通手机照片、AI处理与纸面输出链路。具体功能设计和系统实现将在后续章节展开。",
    )

    insert_paragraph("h2", "设计背景与需求定义")
    insert_paragraph("h3", "数字图像的实体化表达需求")
    insert_paragraph(
        "body",
        "照片便于数字化保存，却往往停留在屏幕浏览中。将照片轮廓转化为真实笔迹，可在数字存储之外增加一种可触摸、可装裱、可赠送的实体表达方式；纸张纹理、颜色和笔具还会赋予作品独特质感与纪念意义。",
    )
    insert_paragraph("h3", "传统写字机的局限与设计目标")
    insert_paragraph(
        "body",
        "典型笛卡尔式XY写字机控制成熟，但常需两组长行程导轨、横梁和外围框架，幅面增大时材料、体积与装配量随之增加；部分产品还依赖专用上位机，内容处理与设备控制相互分离。因此，本作品以低成本、小体积、便携和低使用门槛为主要目标。",
    )

    insert_paragraph("h2", "SCARA结构创新与综合优势")
    insert_paragraph("h3", "以算法复杂度换取机械结构简化")
    insert_paragraph(
        "body",
        "本作品采用平面双关节SCARA机构，以旋转臂覆盖二维书写区域。末端轨迹需通过逆运动学、可达性判断、轨迹插补和误差补偿转换为关节运动，同时对臂长、同轴度、刚度和零位标定提出更严格要求。这一软硬件协同取舍以可复用的软件能力换取长导轨与大框架的减少。",
    )
    insert_paragraph("h3", "低成本、小体积与便携性的形成")
    insert_paragraph(
        "body",
        "双旋转臂有望减少长行程直线模组和专用机械件，并结合通用电机、标准紧固件与模块化结构缩小占用空间、减轻重量和降低收纳难度。成本优势来自材料与结构的优化，而非降低书写质量，最终将结合预算清单和样机实测进行验证。",
    )
    insert_paragraph("h3", "传统写字机与本作品对比")
    insert_paragraph(
        "body",
        "下图从结构、算法和交互三个层面概括两类方案的差异，后续可用同幅面市场产品与本作品样机照片替换占位图。",
    )
    comparison_table = build_comparison_table(doc, image_path)
    section_break.addprevious(comparison_table._tbl)
    insert_paragraph(
        "figure_caption",
        "图1-1 典型笛卡尔式写字机与本作品的结构及体验对比（图片待替换）",
    )

    insert_paragraph("h2", "网页端与AI创作链路创新")
    insert_paragraph("h3", "浏览器统一入口与多源内容输入")
    insert_paragraph(
        "body",
        "网页端作为统一入口，使用户可通过手机、平板或计算机完成文字、插图和照片等内容导入。移动端浏览器在用户授权后可直接调用相机或相册，省去照片传至计算机、再经专用软件导入设备的中间步骤，也便于后续扩展功能。",
    )
    insert_paragraph("h3", "照片线稿化与AI辅助创作")
    insert_paragraph(
        "body",
        "网页端可对照片进行轮廓增强、背景简化和线稿优化，也可依据文字描述生成图案。处理结果结合纸张边界、笔尖宽度和机械精度转换为可执行轨迹，并在下发前完成版面与路径预览，为普通照片增加新的创作可能。",
    )
    insert_paragraph("h3", "从手机照片到纸面作品的闭环")
    insert_paragraph(
        "body",
        "完整链路可概括为“手机拍照或相册选择—网页AI处理与轨迹预览—SCARA纸面输出”。它把内容获取、图像处理和设备执行整合为连续操作，是本作品在使用体验层面的核心特色。",
    )
    workflow_table = build_workflow_table(doc, image_path)
    section_break.addprevious(workflow_table._tbl)
    insert_paragraph(
        "figure_caption",
        "图1-2 手机拍照、网页端处理到写画机纸面输出的闭环流程（图片待替换）",
    )

    insert_paragraph("h2", "纸面体验与创新特色归纳")
    insert_paragraph("h3", "纸张、笔具与成品质感适配")
    insert_paragraph(
        "body",
        "纸面输出不是普通打印复制，而是由真实笔迹、纸张纹理和笔具共同形成的再创作。通过适配不同纸色、纸纹与笔触，同一数字图像可获得差异化的实体质感，更适合展示、装裱、赠送和长期留存。",
    )
    insert_paragraph("h3", "创新特色与后续展开")
    insert_paragraph(
        "body",
        "作品特色可归纳为四点：SCARA结构利于小型化，运动学与标定补偿支撑书写，网页端和AI降低创作门槛，手机照片到纸面形成完整闭环。本章只作概述；功能设计、系统实现与预算清单将在后续章节展开。",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
