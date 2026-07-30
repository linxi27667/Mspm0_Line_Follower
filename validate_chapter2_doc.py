from pathlib import Path
import sys

from docx import Document
from docx.oxml.ns import qn


def paragraph_num(paragraph):
    num_pr = paragraph._p.find("./" + qn("w:pPr") + "/" + qn("w:numPr"))
    if num_pr is None:
        return None, None
    num_id = num_pr.find(qn("w:numId"))
    ilvl = num_pr.find(qn("w:ilvl"))
    return (
        num_id.get(qn("w:val")) if num_id is not None else None,
        ilvl.get(qn("w:val")) if ilvl is not None else None,
    )


def first_line_twips(paragraph):
    indent = paragraph._p.find("./" + qn("w:pPr") + "/" + qn("w:ind"))
    if indent is None:
        return None
    return indent.get(qn("w:firstLine"))


def section_text(doc, start_text):
    start = next(
        index for index, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip() == start_text
    )
    return [paragraph.text for paragraph in doc.paragraphs[start:]]


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: validate_chapter2_doc.py SOURCE.docx OUTPUT.docx")
    source_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    source = Document(source_path)
    doc = Document(output_path)

    start = next(
        index for index, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip() == "特色与创新"
    )
    end = next(
        index for index, paragraph in enumerate(doc.paragraphs)
        if index > start and paragraph.text.strip() == "功能设计"
    )
    chapter = doc.paragraphs[start:end]

    expected_h2 = [
        "低成本与轻量化硬件创新",
        "SCARA运动控制算法创新",
        "网页端、AI与语音交互创新",
        "综合创新与应用优势",
    ]
    actual_h2 = [p.text for p in chapter if p.style.style_id == "2"]
    assert actual_h2 == expected_h2, actual_h2

    for paragraph in chapter:
        if paragraph.style.style_id == "1":
            assert paragraph_num(paragraph) == ("44", "0")
        elif paragraph.style.style_id == "2":
            assert paragraph_num(paragraph) == ("44", "1")
        elif paragraph.style.style_id == "3":
            assert paragraph_num(paragraph) == ("44", "2")

    body_paragraphs = [
        p for p in chapter
        if p.style.style_id == "a"
        and p.text.strip()
        and not p.text.startswith("图2-")
        and not p.text.startswith("表2-")
    ]
    bad_indents = [(p.text[:24], first_line_twips(p)) for p in body_paragraphs
                   if first_line_twips(p) != "480"]
    assert not bad_indents, bad_indents

    figure_captions = [p for p in chapter if p.text.startswith("图2-")]
    assert [p.text[:4] for p in figure_captions] == ["图2-1", "图2-2", "图2-3"]
    for paragraph in figure_captions:
        assert paragraph.alignment is not None
        for run in paragraph.runs:
            if run.text:
                r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
                assert r_fonts.get(qn("w:eastAsia")) == "楷体"
                assert run.font.size is not None and run.font.size.pt == 11.0

    assert any("ESP32-S3" in p.text and "功放" in p.text and "麦克风" in p.text
               for p in chapter)
    assert any("语义分析" in p.text and "结构化命令" in p.text for p in chapter)
    assert any("右图待替换" in p.text for p in chapter)
    assert any("图片待替换" in p.text and "小智" in p.text for p in chapter)

    assert len(doc.sections) == len(source.sections) == 9
    assert len(doc.tables) == len(source.tables) + 3
    assert len(doc.inline_shapes) == len(source.inline_shapes) + 4
    assert section_text(doc, "功能设计") == section_text(source, "功能设计")

    for table in doc.tables[-3:]:
        borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")
        assert borders is not None
        for alternatives in (
            ("top",),
            ("start", "left"),
            ("bottom",),
            ("end", "right"),
            ("insideH",),
            ("insideV",),
        ):
            border = next(
                (borders.find(qn(f"w:{edge}")) for edge in alternatives
                 if borders.find(qn(f"w:{edge}")) is not None),
                None,
            )
            assert border is not None and border.get(qn("w:val")) == "nil"

    print(
        f"OK paragraphs={len(doc.paragraphs)} chapter2_paragraphs={len(chapter)} "
        f"body={len(body_paragraphs)} captions={len(figure_captions)} "
        f"tables={len(doc.tables)} pictures={len(doc.inline_shapes)} sections={len(doc.sections)}"
    )


if __name__ == "__main__":
    main()
