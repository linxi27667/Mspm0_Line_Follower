from copy import deepcopy
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from revise_chapter2_innovation import (
    BODY_FIRST_LINE_TWIPS,
    build_centered_image_paragraph,
    clear_paragraph,
    find_paragraph,
    first_run_properties,
    paragraph_element_from_template,
    set_cell_text,
    set_paragraph_spacing,
    set_run_font,
    set_table_geometry,
)


CHAPTER_START = "功能设计"
CHAPTER_END = "系统实现"


def add_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def build_web_parameter_table(doc):
    widths = [1450, 2140, 2220, 3020]
    table = doc.add_table(rows=7, cols=4)
    set_table_geometry(table, widths)
    headers = ("功能或参数", "用户可调内容", "直接影响", "主要作用")
    for column, text in enumerate(headers):
        set_cell_text(
            table.cell(0, column),
            text,
            font="黑体",
            size=10,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    rows = (
        (
            "一键生成",
            "采用推荐默认值直接处理",
            "快速得到可执行结果",
            "保留低学习成本入口，满足追求效率的用户",
        ),
        (
            "模糊/平滑",
            "调整预处理强度",
            "抑制噪点与零碎纹理",
            "让照片轮廓更适合笔尖表达，减少无效细线",
        ),
        (
            "轨迹简化",
            "设置线段合并或简化程度",
            "改变轮廓细节与轨迹数量",
            "在画面细节、文件规模和书写时间之间自主取舍",
        ),
        (
            "最短轨迹",
            "开启路径排序与空行程优化",
            "减少抬笔后的无效移动",
            "缩短任务时间并降低频繁往返带来的机械消耗",
        ),
        (
            "书写速度",
            "选择速度档位或输入目标速度",
            "改变执行节奏与稳定性",
            "根据笔具、纸张和内容复杂度平衡效率与质量",
        ),
        (
            "幅面与缩放",
            "设置纸张范围、位置和图像尺寸",
            "决定作品在纸面的布局",
            "适配不同纸张与装裱需求，减少越界和重复试画",
        ),
    )
    for row_index, row in enumerate(rows, start=1):
        add_cant_split(table.rows[row_index])
        for column, text in enumerate(row):
            set_cell_text(
                table.cell(row_index, column),
                text,
                font="黑体" if column == 0 else "宋体",
                size=9.2,
                bold=column == 0,
                align=(
                    WD_ALIGN_PARAGRAPH.CENTER
                    if column == 0
                    else WD_ALIGN_PARAGRAPH.JUSTIFY
                ),
            )
    return table


def add_placeholder_image(cell, image_path, width_inches, alt_text):
    paragraph = cell.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, after=0, keep_next=True)
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Inches(width_inches))
    inline._inline.docPr.set("descr", alt_text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def build_paper_output_panel(doc, placeholder_path):
    table = doc.add_table(rows=3, cols=2)
    set_table_geometry(table, [4415, 4415])
    titles = ("写画机纸面成品", "装裱、展示或赠送效果")
    for column, title in enumerate(titles):
        set_cell_text(
            table.cell(0, column),
            title,
            font="黑体",
            size=10.5,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        table.cell(0, column).paragraphs[0].paragraph_format.keep_with_next = True

    specs = (
        (0, "占位图：照片线稿经写画机输出后的纸面成品", "（a）纸面写画成品（待替换）"),
        (1, "占位图：纸面作品装裱展示或赠送场景", "（b）装裱或赠送效果（待替换）"),
    )
    for column, alt_text, subcaption in specs:
        add_placeholder_image(
            table.cell(1, column), placeholder_path, 2.35, alt_text
        )
        set_cell_text(
            table.cell(2, column),
            subcaption,
            font="楷体",
            size=9.5,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for row in table.rows:
        add_cant_split(row)
    return table


def build_ai_workflow(doc, placeholder_path):
    widths = [2600, 515, 2600, 515, 2600]
    table = doc.add_table(rows=2, cols=5)
    set_table_geometry(table, widths)
    specs = (
        (0, "占位图：用户上传的原始照片", "（a）用户原始照片（待替换）", "上传照片或输入提示词"),
        (2, "占位图：AI优化后的线稿预览", "（b）AI优化线稿（待替换）", "生成、修改与参数调整"),
        (4, "占位图：写画机输出的纸面作品", "（c）纸面创作结果（待替换）", "轨迹确认与纸面输出"),
    )
    for column, alt_text, subcaption, label in specs:
        add_placeholder_image(
            table.cell(0, column), placeholder_path, 1.0, alt_text
        )
        note = table.cell(0, column).add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(note, after=0)
        note_run = note.add_run(subcaption)
        set_run_font(note_run, "楷体", 9.2)

        set_cell_text(
            table.cell(1, column),
            label,
            font="宋体",
            size=9.3,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    for column in (1, 3):
        for row_index in (0, 1):
            cell = table.cell(row_index, column)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            clear_paragraph(paragraph)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(paragraph, after=0)
            if row_index == 0:
                run = paragraph.add_run("→")
                set_run_font(run, "宋体", 14, bold=True)
    for row in table.rows:
        add_cant_split(row)
    return table


def build_voice_intent_table(doc):
    widths = [1440, 2730, 2250, 2410]
    table = doc.add_table(rows=5, cols=4)
    set_table_geometry(table, widths)
    headers = ("意图类型", "自然语言示例", "结构化结果", "功能作用")
    for column, text in enumerate(headers):
        set_cell_text(
            table.cell(0, column),
            text,
            font="黑体",
            size=10,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    rows = (
        (
            "文字创作",
            "“帮我写一段生日祝福并写下来”",
            "识别主题，调用AI生成文字并建立书写任务",
            "把内容生成与设备执行连接成一次对话",
        ),
        (
            "图像创作",
            "“把刚拍的照片画得简洁一些”",
            "调用照片、AI线稿化和轨迹预览流程",
            "用自然表达完成跨功能组合操作",
        ),
        (
            "执行控制",
            "“开始写”“先停一下”“停止任务”",
            "映射为开始、暂停或停止等明确指令",
            "支持免手操作，并保留确定性的控制边界",
        ),
        (
            "设备管理",
            "“现在进行到哪里”“让设备归零”",
            "映射为状态查询或归零任务",
            "便于获取反馈并恢复一致的机械基准",
        ),
    )
    for row_index, row in enumerate(rows, start=1):
        add_cant_split(table.rows[row_index])
        for column, text in enumerate(row):
            set_cell_text(
                table.cell(row_index, column),
                text,
                font="黑体" if column == 0 else "宋体",
                size=9.2,
                bold=column == 0,
                align=(
                    WD_ALIGN_PARAGRAPH.CENTER
                    if column == 0
                    else WD_ALIGN_PARAGRAPH.JUSTIFY
                ),
            )
    return table


def main():
    if len(sys.argv) != 4:
        raise SystemExit(
            "Usage: revise_chapter3_function_design.py SOURCE.docx "
            "PLACEHOLDER.jpg OUTPUT.docx"
        )

    source_path = Path(sys.argv[1])
    placeholder_path = Path(sys.argv[2])
    output_path = Path(sys.argv[3])
    if not placeholder_path.exists():
        raise FileNotFoundError(placeholder_path)

    doc = Document(source_path)
    chapter_start = find_paragraph(doc, CHAPTER_START)
    chapter_end = find_paragraph(doc, CHAPTER_END)
    start_index = next(
        index
        for index, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip() == CHAPTER_START
    )
    h2_template = next(
        p for p in doc.paragraphs[start_index + 1 :] if p.style.style_id == "2"
    )
    h3_template = next(
        p for p in doc.paragraphs[start_index + 1 :] if p.style.style_id == "3"
    )
    body_template = find_paragraph(
        doc,
        "本章对照典型笛卡尔式XY写字机与部分市售应用，集中说明本作品在机械结构、运动控制算法、网页端交互、AI创作和语音语义控制方面的特色。这里主要解释创新机制及其形成的优势；具体功能设计和系统实现留待后续章节展开，材料成本则在第五章结合预算清单进行量化验证。",
    )
    figure_caption_template = find_paragraph(
        doc, "图2-1 本作品SCARA写画机的双旋转臂与轻量化结构"
    )
    table_caption_template = find_paragraph(
        doc, "表2-1 典型XY方案与本作品的结构成本构成对比"
    )

    body = doc.element.body
    children = list(body)
    body_start = children.index(chapter_start._p)
    body_end = children.index(chapter_end._p)
    between = children[body_start + 1 : body_end]
    section_break = next(
        (
            element
            for element in reversed(between)
            if element.tag == qn("w:p")
            and element.find("./" + qn("w:pPr") + "/" + qn("w:sectPr"))
            is not None
        ),
        None,
    )
    if section_break is None:
        raise ValueError("Expected a section-break paragraph before the next chapter")
    for element in between:
        if element is not section_break:
            body.remove(element)

    templates = {
        "h2": (h2_template._p, first_run_properties(h2_template)),
        "h3": (h3_template._p, first_run_properties(h3_template)),
        "body": (body_template._p, first_run_properties(body_template)),
        "figure_caption": (
            figure_caption_template._p,
            first_run_properties(figure_caption_template),
        ),
        "table_caption": (
            table_caption_template._p,
            first_run_properties(table_caption_template),
        ),
    }

    def insert_paragraph(kind, text):
        template_p, run_rpr = templates[kind]
        element = paragraph_element_from_template(
            template_p,
            text,
            run_rpr,
            first_line_twips=(
                BODY_FIRST_LINE_TWIPS if kind == "body" else None
            ),
        )
        if kind == "table_caption":
            p_pr = element.get_or_add_pPr()
            if p_pr.find(qn("w:keepNext")) is None:
                p_pr.append(OxmlElement("w:keepNext"))
        section_break.addprevious(element)

    insert_paragraph(
        "body",
        "本章根据前述需求规划作品的主要功能，重点说明各项功能对用户体验和纸面创作流程的作用。整体设计兼顾两类用户：希望快速得到结果的用户可以使用推荐参数一键生成，愿意参与创作的用户则可调整图像、轨迹与速度参数；同时，照片实体化、AI辅助创作和小智语音交互共同构成从内容产生到纸面输出的完整功能链路。",
    )

    insert_paragraph("h2", "网页端双层创作与参数调节功能")
    insert_paragraph("h3", "一键生成与自定义模式并存")
    insert_paragraph(
        "body",
        "部分市售产品以“傻瓜式操作”为卖点，但实际流程往往被压缩为单一的一键生成按钮，用户虽然能够快速得到结果，却难以决定图像应保留多少细节、轨迹如何组织以及设备以何种节奏执行。本作品保留一键生成入口，并提供经过预设的推荐参数，使初次使用者无需学习复杂设置即可完成任务；在此基础上进一步开放自定义模式，把创作选择权交还给用户。",
    )
    insert_paragraph("h3", "图像转换与轨迹参数开放")
    insert_paragraph(
        "body",
        "自定义模式允许用户调节模糊或平滑强度、轨迹简化程度、最短轨迹优化、书写速度以及幅面与缩放等参数。模糊和平滑用于抑制照片噪点，轨迹简化决定轮廓细节和线段数量，最短轨迹减少抬笔后的无效移动，速度设置则帮助用户根据笔具、纸张和内容复杂度平衡效率与稳定性。参数开放使生成结果不再是不可解释的黑箱，而成为用户可以参与和控制的创作过程。",
    )
    insert_paragraph("h3", "即时预览与结果确认")
    insert_paragraph(
        "body",
        "网页端在任务下发前提供图像效果、纸张布局和轨迹预览，使用户能够比较不同参数组合带来的细节变化、线条数量和执行路径，并在确认后再生成正式任务。该功能可减少反复试画造成的纸张和时间浪费，也使一键模式与进阶模式共享同一套确认流程，在便捷性与可控性之间取得平衡。",
    )
    web_image = build_centered_image_paragraph(
        doc,
        placeholder_path,
        4.25,
        "占位图：网页端一键生成、自定义参数与轨迹预览界面",
    )
    section_break.addprevious(web_image._p)
    insert_paragraph(
        "figure_caption",
        "图3-1 网页端一键生成、自定义参数与轨迹预览界面（图片待替换）",
    )
    insert_paragraph("table_caption", "表3-1 网页端主要功能参数及其作用")
    web_table = build_web_parameter_table(doc)
    section_break.addprevious(web_table._tbl)

    insert_paragraph("h2", "照片纸面化与实体留存功能")
    insert_paragraph("h3", "照片导入、线稿转换与纸张适配")
    insert_paragraph(
        "body",
        "用户可通过浏览器上传照片，或在移动端授权后直接调用相机和相册。系统将照片转换为适合笔尖表达的线稿，并依据纸张范围、作品位置和缩放比例生成写画轨迹。与直接打印像素图像不同，该功能关注轮廓、笔触和纸张之间的关系，使同一张照片能够随纸色、纸纹和笔具产生不同的实体表达。",
    )
    insert_paragraph("h3", "展示、装裱、赠送与长期留存")
    insert_paragraph(
        "body",
        "在快节奏生活中，数字照片易于拍摄和分享，却常被淹没在相册和信息流中。本作品提供一种补充性的保存方式：用笔墨和线条重新记录照片，使内容从屏幕文件转化为可触摸的纸面作品。真实笔迹和纸张纹理赋予作品独特质感，更适合摆放展示、装裱收藏、作为礼物赠送或长期留存，也让照片承载更明确的纪念意义。",
    )
    paper_panel = build_paper_output_panel(doc, placeholder_path)
    section_break.addprevious(paper_panel._tbl)
    insert_paragraph(
        "figure_caption",
        "图3-2 照片线稿的纸面成品及装裱应用效果（图片待替换）",
    )

    insert_paragraph("h2", "AI辅助创作与再创作功能")
    insert_paragraph("h3", "文字与图像内容生成")
    insert_paragraph(
        "body",
        "当用户只有主题或想法而缺少现成素材时，可通过提示词让AI生成文字、祝福语、短句或图案。生成结果并不直接绕过用户进入设备，而是先回到网页端进行查看和修改，再结合纸张版面与轨迹参数形成任务。该功能降低了内容准备门槛，使不会绘画或暂时缺少创作灵感的用户也能获得可进一步编辑的起点。",
    )
    insert_paragraph("h3", "上传照片优化与人机共创")
    insert_paragraph(
        "body",
        "对于用户上传的照片，AI可辅助完成主体提取、背景简化、轮廓增强、线稿化和风格调整，为普通照片增加新的创作可能。用户仍可继续调节模糊、轨迹简化和路径优化等参数，并选择是否采用生成结果。AI负责提供候选方案，用户负责判断和取舍，最终输出既保留原始照片信息，也体现个人审美和纸面表现需求。",
    )
    insert_paragraph("h3", "从AI结果到可执行纸面轨迹")
    insert_paragraph(
        "body",
        "AI生成或优化后的内容需要经过版面适配、线条筛选和轨迹预览，才能转化为写画机能够执行的任务。该功能把AI内容生产与真实机械输出连接起来，并通过人工确认避免不可写、过密或超出纸张范围的结果直接进入设备，使AI的开放式创作能力落到可控的纸面作品中。",
    )
    ai_workflow = build_ai_workflow(doc, placeholder_path)
    section_break.addprevious(ai_workflow._tbl)
    insert_paragraph(
        "figure_caption",
        "图3-3 原始照片、AI线稿优化与写画机纸面输出流程（图片待替换）",
    )

    insert_paragraph("h2", "小智语音语义控制功能")
    insert_paragraph("h3", "自然语言输入与上下文理解")
    insert_paragraph(
        "body",
        "小智通过ESP32-S3语音单元和麦克风接收用户表达，并结合上下文理解其真实意图。同一操作无需依赖固定口令，例如“开始写吧”“可以动笔了”或“执行刚才的内容”均可被理解为开始任务；连续对话还可承接前文中的照片、文字主题或参数要求。与传统语音遥控相比，该功能减少了用户记忆命令词的负担，使交互更接近日常交流。",
    )
    insert_paragraph("h3", "意图转换与跨功能联动")
    insert_paragraph(
        "body",
        "小智不仅控制开始和停止，还可把语义意图连接到AI文字生成、图片处理、网页任务创建和设备状态查询。例如，用户提出“把刚拍的照片画得简洁一些”时，系统需要识别照片对象、简化要求和绘图目标，再进入AI线稿化与轨迹预览流程。语音入口由此成为多项功能的统一调度方式，而不是只能执行单个动作的遥控按钮。",
    )
    insert_paragraph("table_caption", "表3-2 小智语音意图与功能联动示例")
    voice_table = build_voice_intent_table(doc)
    section_break.addprevious(voice_table._tbl)
    insert_paragraph("h3", "结构化命令与安全确认")
    insert_paragraph(
        "body",
        "自然语言具有省略、歧义和表达多样性，不能直接驱动机械执行。小智先将语义转换为明确的任务类型、对象和参数，再由结构化命令层检查必要信息、设备状态和动作条件；涉及开始、停止、归零或参数缺失的任务，可在执行前进行确认或反馈。该设计隔离了自然语言的不确定性，在提升免手操作便捷性的同时，保留写画机控制所需的确定性和可靠性。",
    )
    insert_paragraph(
        "body",
        "四组功能共同形成“快捷入口与自主调节并存、数字内容与纸面作品互补、AI生成与人工取舍结合、自然语言与结构化控制衔接”的使用体系。第三章说明的是功能规划及其作用，相关网页架构、AI服务、通信链路和终端控制实现将在第四章进一步展开。",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
