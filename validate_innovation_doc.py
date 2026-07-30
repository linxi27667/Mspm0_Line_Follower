from pathlib import Path
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn


def find_index(doc, text):
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == text:
            return index
    raise AssertionError(f"Missing paragraph: {text}")


def numbering(paragraph):
    p_pr = paragraph._p.pPr
    if p_pr is None or p_pr.numPr is None:
        return None
    return int(p_pr.numPr.numId.val), int(p_pr.numPr.ilvl.val)


def signature(paragraph):
    return paragraph.text, paragraph.style.style_id, numbering(paragraph)


def table_geometry(table):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    grid = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
    cell_widths = []
    for cell in table.rows[0].cells:
        tc_w = cell._tc.tcPr.tcW
        cell_widths.append(int(tc_w.get(qn("w:w"))))
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    border_values = {}
    for edge, variants in {
        "top": ("top",),
        "start": ("start", "left"),
        "bottom": ("bottom",),
        "end": ("end", "right"),
        "insideH": ("insideH",),
        "insideV": ("insideV",),
    }.items():
        node = next(
            (borders.find(qn(f"w:{variant}")) for variant in variants if borders.find(qn(f"w:{variant}")) is not None),
            None,
        )
        border_values[edge] = None if node is None else node.get(qn("w:val"))
    return (
        int(tbl_w.get(qn("w:w"))),
        int(tbl_ind.get(qn("w:w"))),
        grid,
        cell_widths,
        border_values,
    )


def main():
    source_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])

    with zipfile.ZipFile(output_path) as package:
        assert package.testzip() is None
        assert "word/document.xml" in package.namelist()

    source = Document(source_path)
    output = Document(output_path)
    source_start = find_index(source, "设计需求分析")
    output_start = find_index(output, "设计需求分析")
    source_end = find_index(source, "特色与创新")
    output_end = find_index(output, "特色与创新")

    assert [signature(p) for p in source.paragraphs[: source_start + 1]] == [
        signature(p) for p in output.paragraphs[: output_start + 1]
    ]
    assert [signature(p) for p in source.paragraphs[source_end:]] == [
        signature(p) for p in output.paragraphs[output_end:]
    ]
    assert len(source.sections) == len(output.sections) == 9

    chapter = output.paragraphs[output_start + 1 : output_end]
    h2 = [p for p in chapter if p.style.style_id == "2"]
    h3 = [p for p in chapter if p.style.style_id == "3"]
    assert len(h2) == 4
    assert len(h3) == 10
    assert all(numbering(p) == (44, 1) for p in h2)
    assert all(numbering(p) == (44, 2) for p in h3)
    assert any(p.text == "数字图像的实体化表达需求" for p in h3)
    assert any(p.text == "照片线稿化与AI辅助创作" for p in h3)
    assert any(p.text == "纸张、笔具与成品质感适配" for p in h3)
    assert any(p.text == "创新特色与后续展开" for p in h3)

    captions = [p for p in chapter if p.text.startswith("图1-")]
    assert [p.text[:4] for p in captions] == ["图1-1", "图1-2"]
    for caption in captions:
        assert caption.alignment is not None
        first_run = next(run for run in caption.runs if run.text)
        assert first_run.font.size.pt == 11
        assert first_run._element.rPr.rFonts.get(qn("w:eastAsia")) == "楷体"

    assert len(output.tables) == 3
    comparison = output.tables[1]
    workflow = output.tables[2]
    comparison_geometry = table_geometry(comparison)
    workflow_geometry = table_geometry(workflow)
    assert comparison_geometry[:4] == (8830, 120, [4415, 4415], [4415, 4415])
    assert workflow_geometry[:4] == (
        8830,
        120,
        [2600, 515, 2600, 515, 2600],
        [2600, 515, 2600, 515, 2600],
    )
    assert set(comparison_geometry[4].values()) == {"nil"}
    assert set(workflow_geometry[4].values()) == {"nil"}
    assert len(output.inline_shapes) == 5

    doc_pr_nodes = output.element.body.findall(".//" + qn("wp:docPr"))
    alt_texts = [node.get("descr", "") for node in doc_pr_nodes]
    assert len([text for text in alt_texts if text.startswith("占位图")]) == 5

    print(
        f"ok paragraphs={len(output.paragraphs)} sections={len(output.sections)} "
        f"tables={len(output.tables)} images={len(output.inline_shapes)} "
        f"chapter_h2={len(h2)} chapter_h3={len(h3)} captions={len(captions)}"
    )


if __name__ == "__main__":
    main()
