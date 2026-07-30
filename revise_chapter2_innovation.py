from copy import deepcopy
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


CHAPTER_START = "特色与创新"
CHAPTER_END = "功能设计"
BODY_FIRST_LINE_TWIPS = 480


def find_paragraph(doc, text, start=0):
    for paragraph in doc.paragraphs[start:]:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def first_run_properties(paragraph):
    for run in paragraph.runs:
        if run._r.rPr is not None:
            return run._r.rPr
    return None


def set_first_line_indent(paragraph_element, twips=BODY_FIRST_LINE_TWIPS):
    p_pr = paragraph_element.get_or_add_pPr()
    indent = p_pr.find(qn("w:ind"))
    if indent is None:
        indent = OxmlElement("w:ind")
        p_pr.append(indent)
    for attr in ("hanging", "hangingChars", "firstLineChars"):
        indent.attrib.pop(qn(f"w:{attr}"), None)
    indent.set(qn("w:firstLine"), str(twips))


def paragraph_element_from_template(
    template_p, text, run_rpr=None, first_line_twips=None
):
    paragraph = deepcopy(template_p)
    for child in list(paragraph):
        if child.tag != qn("w:pPr"):
            paragraph.remove(child)
    if first_line_twips is not None:
        set_first_line_indent(paragraph, first_line_twips)

    run = OxmlElement("w:r")
    if run_rpr is not None:
        run.append(deepcopy(run_rpr))
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    paragraph.append(run)
    return paragraph


def set_run_font(run, east_asia, size, bold=None, color=None, italic=None):
    run.font.name = east_asia
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.0, keep_next=None):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if keep_next is not None:
        paragraph.paragraph_format.keep_with_next = keep_next


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_cell_margins(cell, top=75, start=105, bottom=75, end=105):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = None
    total_width = sum(widths)

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total_width))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "nil")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_cell_text(
    cell,
    text,
    *,
    font="宋体",
    size=9.5,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    color=None,
):
    paragraph = cell.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = align
    set_paragraph_spacing(paragraph, after=0, line=1.05)
    run = paragraph.add_run(text)
    set_run_font(run, font, size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def build_cost_table(doc):
    widths = [1300, 2550, 2550, 2430]
    table = doc.add_table(rows=5, cols=4)
    set_table_geometry(table, widths)
    headers = ("对比维度", "典型笛卡尔式XY方案", "本作品SCARA方案", "形成的优势")
    for column, text in enumerate(headers):
        set_cell_text(
            table.cell(0, column),
            text,
            font="黑体",
            size=10,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    rows = (
        (
            "承载结构",
            "两组长导轨、滑块、横梁及型材或钢制框架",
            "两根金属管材构成旋转臂，关节座和笔夹以塑料件或3D打印件为主",
            "减少长直线构件和金属框架用量",
        ),
        (
            "传动部件",
            "通常配置两套直线传动模组及相应张紧、连接件",
            "以双旋转关节覆盖二维书写区域，专用机械件较少",
            "降低物料、装配及维护复杂度",
        ),
        (
            "体积与收纳",
            "工作幅面扩大时，横梁和外围框架同步增大",
            "旋转臂围绕基座展开，非工作状态更易收拢",
            "有利于小型化、轻量化和便携部署",
        ),
        (
            "复杂度所在",
            "坐标映射直观，机械用料与几何约束相对明显",
            "正逆运动学、轨迹规划、标定及局部刚度要求更高",
            "以可复用的软件能力换取机械结构简化",
        ),
    )
    for row_index, row in enumerate(rows, start=1):
        for column, text in enumerate(row):
            set_cell_text(
                table.cell(row_index, column),
                text,
                font="黑体" if column == 0 else "宋体",
                size=9.2,
                bold=column == 0,
                align=WD_ALIGN_PARAGRAPH.CENTER if column == 0 else WD_ALIGN_PARAGRAPH.JUSTIFY,
            )
    return table


def build_input_modes_table(doc):
    widths = [1350, 2230, 2650, 2600]
    table = doc.add_table(rows=4, cols=4)
    set_table_geometry(table, widths)
    headers = ("内容入口", "输入方式", "网页端处理", "输出优势")
    for column, text in enumerate(headers):
        set_cell_text(
            table.cell(0, column),
            text,
            font="黑体",
            size=10,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    rows = (
        (
            "手动文字",
            "直接输入标题、段落或短句",
            "AI生成、润色或改写，结合版面参数预览",
            "兼顾自主编辑与快速创作，降低文字准备门槛",
        ),
        (
            "文档导入",
            "上传已有文档并提取文字与段落",
            "整理内容层级、辅助排版并转换为书写任务",
            "减少重复录入，便于处理长篇或既有材料",
        ),
        (
            "图片与AI创作",
            "上传照片、插图，或输入图像提示词",
            "生成图片、修改上传图片、线稿化并预览轨迹",
            "使写字机扩展为可输出图文内容的写画终端",
        ),
    )
    for row_index, row in enumerate(rows, start=1):
        for column, text in enumerate(row):
            set_cell_text(
                table.cell(row_index, column),
                text,
                font="黑体" if column == 0 else "宋体",
                size=9.2,
                bold=column == 0,
                align=WD_ALIGN_PARAGRAPH.CENTER if column == 0 else WD_ALIGN_PARAGRAPH.JUSTIFY,
            )
    return table


def build_interface_comparison(doc, market_image, placeholder_image):
    table = doc.add_table(rows=3, cols=2)
    set_table_geometry(table, [4415, 4415])
    titles = ("部分市售产品的软件界面", "本作品的网页端统一入口")
    for column, title in enumerate(titles):
        set_cell_text(
            table.cell(0, column),
            title,
            font="黑体",
            size=10.5,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        table.cell(0, column).paragraphs[0].paragraph_format.keep_with_next = True

    specs = (
        (
            0,
            market_image,
            "部分市售产品宣传图，展示多平台客户端和文字抄写界面",
            "（a）市售产品界面示例",
        ),
        (
            1,
            placeholder_image,
            "占位图：本作品网页端三类内容入口",
            "（b）本作品网页端截图（待替换）",
        ),
    )
    for column, path, alt_text, subcaption in specs:
        image_cell = table.cell(1, column)
        paragraph = image_cell.paragraphs[0]
        clear_paragraph(paragraph)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(paragraph, after=0, keep_next=True)
        run = paragraph.add_run()
        inline = run.add_picture(str(path), width=Inches(2.35))
        inline._inline.docPr.set("descr", alt_text)
        image_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        caption_cell = table.cell(2, column)
        set_cell_text(
            caption_cell,
            subcaption,
            font="楷体",
            size=9.5,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    return table


def build_centered_image_paragraph(doc, image_path, width_inches, alt_text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, after=0, keep_next=True)
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Inches(width_inches))
    inline._inline.docPr.set("descr", alt_text)
    return paragraph


def main():
    if len(sys.argv) != 7:
        raise SystemExit(
            "Usage: revise_chapter2_innovation.py SOURCE.docx TEMPLATE.docx "
            "HARDWARE.jpg MARKET.jpg PLACEHOLDER.jpg OUTPUT.docx"
        )

    source_path = Path(sys.argv[1])
    template_path = Path(sys.argv[2])
    hardware_path = Path(sys.argv[3])
    market_path = Path(sys.argv[4])
    placeholder_path = Path(sys.argv[5])
    output_path = Path(sys.argv[6])
    for image_path in (hardware_path, market_path, placeholder_path):
        if not image_path.exists():
            raise FileNotFoundError(image_path)

    doc = Document(source_path)
    original_template = Document(template_path)
    chapter_start = find_paragraph(doc, CHAPTER_START)
    chapter_end = find_paragraph(doc, CHAPTER_END)

    start_index = next(
        index
        for index, paragraph in enumerate(doc.paragraphs)
        if paragraph.text.strip() == CHAPTER_START
    )
    h2_template = next(
        p for p in doc.paragraphs[start_index + 1 :] if p.style.style_id == "2"
    )
    h3_template = next(
        p for p in doc.paragraphs[start_index + 1 :] if p.style.style_id == "3"
    )
    body_template = next(
        p
        for p in doc.paragraphs[start_index + 1 :]
        if p.style.style_id == "a" and p.text.strip()
    )
    figure_caption_template = find_paragraph(
        doc, "图1-1 典型笛卡尔式写字机与本作品的结构及体验对比（图片待替换）"
    )
    table_caption_template = find_paragraph(
        original_template, "表1-1云计算同传统IT服务模式的区别"
    )

    body = doc.element.body
    children = list(body)
    body_start = children.index(chapter_start._p)
    body_end = children.index(chapter_end._p)
    between = children[body_start + 1 : body_end]
    section_break = next(
        (
            element
            for element in reversed(between)
            if element.tag == qn("w:p")
            and element.find("./" + qn("w:pPr") + "/" + qn("w:sectPr"))
            is not None
        ),
        None,
    )
    if section_break is None:
        raise ValueError("Expected a section-break paragraph before the next chapter")
    for element in between:
        if element is not section_break:
            body.remove(element)

    templates = {
        "h2": (h2_template._p, first_run_properties(h2_template)),
        "h3": (h3_template._p, first_run_properties(h3_template)),
        "body": (body_template._p, first_run_properties(body_template)),
        "figure_caption": (
            figure_caption_template._p,
            first_run_properties(figure_caption_template),
        ),
        "table_caption": (
            table_caption_template._p,
            first_run_properties(table_caption_template),
        ),
    }

    def insert_paragraph(kind, text):
        template_p, run_rpr = templates[kind]
        first_line = BODY_FIRST_LINE_TWIPS if kind == "body" else None
        element = paragraph_element_from_template(
            template_p, text, run_rpr, first_line_twips=first_line
        )
        section_break.addprevious(element)

    insert_paragraph(
        "body",
        "本章对照典型笛卡尔式XY写字机与部分市售应用，集中说明本作品在机械结构、运动控制算法、网页端交互、AI创作和语音语义控制方面的特色。这里主要解释创新机制及其形成的优势；具体功能设计和系统实现留待后续章节展开，材料成本则在第五章结合预算清单进行量化验证。",
    )

    insert_paragraph("h2", "低成本与轻量化硬件创新")
    insert_paragraph("h3", "典型笛卡尔式写字机的成本来源")
    insert_paragraph(
        "body",
        "典型笛卡尔式XY写字机通常依靠两组长行程导轨、滑块、同步带或丝杆完成平面定位，并配置横梁、型材或钢制框架保证刚度。该结构坐标关系直观、控制方案成熟，但当书写幅面扩大时，导轨长度、框架尺寸、连接件数量和装配工作量往往同步增加，材料成本、整机重量与收纳空间也随之上升。",
    )
    insert_paragraph("h3", "SCARA轻量化结构方案")
    insert_paragraph(
        "body",
        "本作品采用平面双关节SCARA机构，以两根金属管材构成主要旋转臂，关节座、笔夹和部分连接件采用塑料件或3D打印件。双臂围绕基座转动即可覆盖二维书写区域，减少了长直线导轨、大型横梁和外围框架的使用；在非工作状态下，旋转臂也更易收拢，从而为低成本、小体积、轻量化和便携部署提供结构基础。",
    )
    hardware_image = build_centered_image_paragraph(
        doc,
        hardware_path,
        4.75,
        "本作品SCARA写画机样机实拍，展示双旋转臂和轻量化结构",
    )
    section_break.addprevious(hardware_image._p)
    insert_paragraph(
        "figure_caption",
        "图2-1 本作品SCARA写画机的双旋转臂与轻量化结构",
    )
    insert_paragraph(
        "table_caption", "表2-1 典型XY方案与本作品的结构成本构成对比"
    )
    cost_table = build_cost_table(doc)
    section_break.addprevious(cost_table._tbl)
    insert_paragraph("h3", "成本优势的形成与工程代价")
    insert_paragraph(
        "body",
        "机械用料减少并不意味着设计要求降低，而是把一部分复杂度从机械框架转移到运动学算法、标定补偿和局部结构精度上。为保证笔尖轨迹准确，臂长、关节同轴度、连接刚度和零位一致性需要严格控制。其核心创新不是简单减料，而是以可重复使用的软件与标定能力换取专用机械件的减少，并在保持书写能力的前提下降低综合成本。",
    )

    insert_paragraph("h2", "SCARA运动控制算法创新")
    insert_paragraph("h3", "三角几何建模与正逆运动学")
    insert_paragraph(
        "body",
        "笛卡尔式机构的X、Y坐标可较直接地映射到两条直线轴，而SCARA笔尖位置由两段臂长和两个关节角共同决定。本作品依据三角几何关系建立正、逆运动学模型，通过余弦关系和atan2函数求解关节角，并处理正肘、负肘两种构型；计算后还进行工作区可达性、关节角范围和正解回代误差检查，避免不可达坐标或错误构型进入执行环节。",
    )
    insert_paragraph("h3", "轨迹规划、同步控制与安全约束")
    insert_paragraph(
        "body",
        "为把文字和图像轮廓稳定地落到纸面，系统对G代码直线轨迹进行分段插补和进给速度处理，再将每个末端坐标转换为双关节运动。底层采用双轴同步DDA，使两个步进轴按比例同时到达目标位置，并结合AMASS低速平滑、拐角限速以及前向和反向加减速规划，减轻低速抖动和关节急停冲击。可达区间、角度限制和任务状态检查共同构成执行前的安全约束。",
    )
    insert_paragraph("h3", "归零标定与软硬件协同")
    insert_paragraph(
        "body",
        "SCARA结构的误差会沿连杆传递到笔尖，因此系统设计了限位释放、归零和关节零位标定流程，使控制器建立一致的机械坐标基准。运动学求解、轨迹规划和标定流程与轻量化机械结构协同工作：硬件越简化，算法和装配标定越需要精确；这种复杂度转移正是本作品区别于典型XY方案的技术取舍。",
    )

    insert_paragraph("h2", "网页端、AI与语音交互创新")
    insert_paragraph("h3", "从单一抄写到多源内容创作")
    insert_paragraph(
        "body",
        "部分市售写字机已支持计算机、平板或手机连接，但其功能边界主要集中于文字抄写、字迹复现和设备参数设置。本作品将网页端作为统一创作入口，在保留文字书写能力的同时，引入文档导入、图片处理和AI生成，使设备从执行既有文本扩展为面向文字、文档、照片和图案的图文创作终端。",
    )
    interface_table = build_interface_comparison(doc, market_path, placeholder_path)
    section_break.addprevious(interface_table._tbl)
    insert_paragraph(
        "figure_caption",
        "图2-2 部分市售产品界面与本作品网页端创作入口对比（右图待替换）",
    )
    insert_paragraph("h3", "三类内容输入的统一设计")
    insert_paragraph(
        "body",
        "网页端把内容准备归纳为手动文字、文档导入以及图片与AI创作三类入口。用户无需在多个独立软件之间反复转换文件，便可在同一流程中完成内容输入、AI辅助处理、版面与轨迹预览，再生成终端可执行的书写或绘画任务。",
    )
    insert_paragraph("table_caption", "表2-2 网页端三类内容输入模式及输出优势")
    input_table = build_input_modes_table(doc)
    section_break.addprevious(input_table._tbl)
    insert_paragraph("h3", "AI生成与上传图片再创作")
    insert_paragraph(
        "body",
        "AI能力既可依据提示词生成文字或图案，也可对用户上传的照片进行轮廓增强、背景简化、线稿化和风格调整。处理后的结果再结合纸张边界、笔尖宽度和机械精度转化为可执行轨迹。这样，用户不仅能够抄写已有内容，还能把普通照片和自然语言描述转化为具有真实笔迹、纸张纹理和个性化笔触的实体作品。",
    )
    insert_paragraph("h3", "小智语音助手的语义控制")
    insert_paragraph(
        "body",
        "小智语音助手搭载于ESP32-S3平台。围绕语音交互需求，作品为其设计了专用拓展板，在板上集成麦克风采集与功放输出电路，使语音输入、处理通信和声音反馈形成相对完整的硬件单元。相比直接拼接多个外部模块，自研拓展板有利于减少接线和占用空间，提高系统集成度，也便于与写字机主体统一安装和调试。",
    )
    insert_paragraph(
        "body",
        "用户可以使用自然语言表达需求，而不必逐级点击固定菜单。小智先完成语音内容识别和语义分析，再把“写一段祝福语”“绘制指定图片”“停止当前任务”“设备归零”或“查询状态”等意图映射为写字机能够识别的结构化命令。写字机端接收明确的任务类型和参数后再执行相应动作，从而把开放式语言交互与确定性的机械控制连接起来。",
    )
    insert_paragraph(
        "body",
        "这种设计的创新点在于，小智不仅充当语音遥控器，还承担上下文理解和意图转换，使同一操作可以用不同说法触发，并可与AI文字生成、图片处理及网页任务流程衔接。结构化命令层同时隔离自然语言的不确定性，便于在执行前进行参数和设备状态检查，提升免手操作的便捷性与控制可靠性。",
    )
    voice_image = build_centered_image_paragraph(
        doc,
        placeholder_path,
        3.25,
        "占位图：ESP32-S3小智语音拓展板及其与写字机的语义控制联动",
    )
    section_break.addprevious(voice_image._p)
    insert_paragraph(
        "figure_caption",
        "图2-3 基于ESP32-S3的小智语音拓展板及语义控制联动场景（图片待替换）",
    )
    insert_paragraph("h3", "浏览器跨端与端云闭环")
    insert_paragraph(
        "body",
        "采用浏览器入口后，手机、平板和计算机可共享同一交互方式；移动端在用户授权后还可直接调用相机或相册，缩短照片从拍摄到处理的路径。系统按浏览器、服务端和终端控制器分层协同，终端侧已具备基于ESP-AT与MQTT的任务接收、状态上报和保活基础，可形成“手机拍照或选图—网页AI处理—轨迹预览与任务下发—写画机纸面输出”的连续链路。",
    )

    insert_paragraph("h2", "综合创新与应用优势")
    insert_paragraph("h3", "从抄写设备到图文创作终端")
    insert_paragraph(
        "body",
        "本作品的创新并非单一部件替换，而是结构、算法、交互与内容生产方式的组合：SCARA结构减少机械用料，正逆运动学和轨迹规划支撑准确书写，网页端统一多源内容入口，AI拓展文字和图片的生成与再创作，小智语音助手把自然语言意图转换为设备命令，最终以真实笔迹完成纸面表达。由此形成低成本、小型便携、交互方式自然和创作形式丰富的综合优势。",
    )
    insert_paragraph("h3", "分层架构与扩展价值")
    insert_paragraph(
        "body",
        "网页端、服务端、通信链路和终端控制器相互分层，使内容处理能力与机械执行能力可以分别迭代。后续既可扩展新的AI模型、文档解析和图像处理方式，也可复用同一运动控制与任务通信框架适配不同纸张、笔具和书写幅面。具体功能流程、软硬件实现以及成本清单将在第三章至第五章继续说明。",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
